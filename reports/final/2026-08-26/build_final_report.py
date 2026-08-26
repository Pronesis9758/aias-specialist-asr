"""Build the evidence-backed AI Specialist manufacturing ASR final report."""

from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
OUT_DIR = Path(__file__).resolve().parent
ASSET_DIR = OUT_DIR / "assets"
EVIDENCE_PATH = OUT_DIR / "evidence" / "final_report_evidence.json"
DOCX_PATH = OUT_DIR / "AI_Specialist_ASR_최종기술보고서_일반화검증반영_2026-08-26.docx"
EVIDENCE = json.loads(EVIDENCE_PATH.read_text(encoding="utf-8"))

NAVY = "17365D"
BLUE = "2E74B5"
MID_BLUE = "5B9BD5"
PALE_BLUE = "DDEBF7"
VERY_PALE_BLUE = "F3F7FB"
GREEN = "2E7D32"
PALE_GREEN = "E8F5E9"
ORANGE = "C55A11"
PALE_ORANGE = "FCE4D6"
RED = "C00000"
PALE_RED = "FDE9E7"
DARK = "243447"
GRAY = "667085"
LIGHT_GRAY = "E6E9ED"
WHITE = "FFFFFF"


def pct(value: float, digits: int = 2) -> str:
    return f"{value * 100:.{digits}f}%"


def gb(size_bytes: int) -> str:
    return f"{size_bytes / 1_000_000_000:.2f} GB"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text: str, *, bold=False, color=DARK, size=8.5, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[object]], widths=None):
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        set_cell_shading(hdr.cells[idx], NAVY)
        set_cell_text(hdr.cells[idx], header, bold=True, color=WHITE, size=8.1, align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            hdr.cells[idx].width = Cm(widths[idx])
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        set_row_no_split(row)
        for col_idx, value in enumerate(values):
            if row_idx % 2 == 1:
                set_cell_shading(row.cells[col_idx], VERY_PALE_BLUE)
            align = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(row.cells[col_idx], str(value), size=7.8, align=align)
            if widths:
                row.cells[col_idx].width = Cm(widths[col_idx])
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("AI Specialist 제조 ASR  ·  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.13

    for style_name, size, color in (("Title", 28, NAVY), ("Heading 1", 18, NAVY), ("Heading 2", 12.5, BLUE), ("Heading 3", 10.5, DARK)):
        style = doc.styles[style_name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = False
        style.paragraph_format.space_before = Pt(10 if style_name != "Heading 1" else 14)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header.add_run("제조 현장 한국어 ASR 성능·일반화·온디바이스 연구")
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    add_page_number(section.footer.paragraphs[0])


def add_body(doc: Document, text: str, *, bold=False, color=DARK, size=9.5, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.2)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        r.font.name = "Malgun Gothic"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        r.font.size = Pt(9.2)


def add_callout(doc: Document, title: str, text: str, *, kind="info") -> None:
    palette = {
        "info": (PALE_BLUE, NAVY),
        "pass": (PALE_GREEN, GREEN),
        "warn": (PALE_ORANGE, ORANGE),
        "fail": (PALE_RED, RED),
    }
    fill, color = palette[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_row_no_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    r.font.size = Pt(10.3)
    r.font.color.rgb = RGBColor.from_string(color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(text)
    r2.font.name = "Malgun Gothic"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    r2.font.size = Pt(9.2)
    r2.font.color.rgb = RGBColor.from_string(DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GRAY)


def add_section(doc: Document, title: str, *, page_break=True) -> None:
    heading = doc.add_heading(title, level=1)
    # A paragraph-level page break avoids the extra blank page that can occur
    # when an explicit break is pushed to the next page by a full preceding page.
    if page_break:
        heading.paragraph_format.page_break_before = True


def font(size: int, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/malgun.ttf"),
        Path("C:/Windows/Fonts/malgunbd.ttf") if bold else Path("C:/Windows/Fonts/malgun.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def canvas(title: str, subtitle: str, height=850):
    image = Image.new("RGB", (1600, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1600, 12), fill="#2E74B5")
    draw.text((70, 45), title, font=font(34, True), fill="#17365D")
    draw.text((70, 98), subtitle, font=font(18), fill="#667085")
    draw.line((70, 142, 1530, 142), fill="#D0D5DD", width=2)
    return image, draw


def draw_centered(draw, box, text, *, size=22, fill="#243447", bold=False):
    x0, y0, x1, y1 = box
    fnt = font(size, bold)
    bbox = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=6, align="center")
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    draw.multiline_text(((x0 + x1 - width) / 2, (y0 + y1 - height) / 2), text, font=fnt, fill=fill, spacing=6, align="center")


def draw_arrow(draw, start, end, color="#98A2B3", width=5):
    draw.line((start, end), fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 16
    for offset in (2.6, -2.6):
        point = (end[0] + length * math.cos(angle + offset), end[1] + length * math.sin(angle + offset))
        draw.line((end, point), fill=color, width=width)


def create_charts() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    image, draw = canvas("연구 전체 흐름", "Test를 보기 전에 Validation에서 모든 선택을 끝내는 단방향 프로토콜", 520)
    labels = ["데이터\n무결성", "일반 모델\nBenchmark", "LoRA\n비교", "Decoding·IR/NN\nValidation", "양자화\n비교", "Frozen Test\nv1·v2·v3"]
    colors = ["#DDEBF7", "#E8F5E9", "#FFF2CC", "#E4DFEC", "#FCE4D6", "#FDE9E7"]
    x = 70
    boxes = []
    box_widths = [210, 210, 210, 225, 200, 220]
    for idx, (label, color) in enumerate(zip(labels, colors)):
        width = box_widths[idx]
        box = (x, 210, x + width, 355)
        draw.rounded_rectangle(box, radius=16, fill=color, outline="#B8C2CC", width=2)
        draw_centered(draw, box, label, size=21, bold=True)
        boxes.append(box)
        x += width + 20
    for left, right in zip(boxes, boxes[1:]):
        draw_arrow(draw, (left[2] + 4, 282), (right[0] - 5, 282))
    draw.text((70, 410), "선택 영역: Validation only", font=font(19, True), fill="#2E7D32")
    draw.text((700, 410), "보고 영역: Test 결과를 보고 같은 Test에 재적합하지 않음", font=font(19, True), fill="#C00000")
    path = ASSET_DIR / "research_flow_simple.png"
    image.save(path)
    paths["flow"] = path

    image, draw = canvas("재현 가능한 ASR 연구 아키텍처", "데이터 → AIAS CLI → GPU 학습·추론 → 불변 산출물·심사 증빙", 990)
    layers = [
        ("DATA & GOVERNANCE", 180, ["manifest·provenance", "speaker/text split", "합성·승인 상태", "hash·중복 검사"]),
        ("AIAS PIPELINE", 370, ["prepare", "benchmark", "LoRA train", "decode sweep", "IR/NN gate", "quantize", "validate"]),
        ("COMPUTE", 585, ["Colab A100 80GB", "고용량 RAM", "HF pinned model", "CTranslate2"]),
        ("EVIDENCE", 790, ["config snapshot", "predictions", "metrics·CI", "SQLite registry", "DOCX report"]),
    ]
    palette = ["#DDEBF7", "#E8F5E9", "#FFF2CC", "#FCE4D6"]
    for layer_idx, (name, y, items) in enumerate(layers):
        draw.text((70, y), name, font=font(20, True), fill="#17365D")
        box_y0, box_y1 = y + 42, y + 142
        available = 1450
        gap = 14
        width = (available - gap * (len(items) - 1)) / len(items)
        for idx, item in enumerate(items):
            x0 = 70 + idx * (width + gap)
            box = (x0, box_y0, x0 + width, box_y1)
            draw.rounded_rectangle(box, radius=12, fill=palette[layer_idx], outline="#B8C2CC", width=2)
            draw_centered(draw, box, item, size=18, bold=True)
        if layer_idx < len(layers) - 1:
            draw_arrow(draw, (800, box_y1 + 5), (800, layers[layer_idx + 1][1] - 8), width=4)
    path = ASSET_DIR / "research_architecture_detailed.png"
    image.save(path)
    paths["architecture"] = path

    legacy = EVIDENCE["legacy_study"]["dataset"]
    general = EVIDENCE["generalization_study"]["development_dataset"]
    image, draw = canvas("데이터 설계 전후", "기존 성능 입증 데이터는 보존하고 일반화 개발·독립 Test를 추가", 760)
    rows = [
        ("Legacy Train", legacy["train_samples"], "#5B9BD5"),
        ("Legacy Validation", legacy["validation_samples"], "#9CC3E5"),
        ("Legacy Test v1", legacy["test_samples"], "#B4C6E7"),
        ("Generalization Train", general["train_samples"], "#70AD47"),
        ("Generalization Validation", general["validation_samples"], "#A9D18E"),
        ("Independent Test v3", EVIDENCE["generalization_study"]["independent_test_v3"]["samples"], "#F4B183"),
    ]
    x0, x1, max_value = 470, 1470, 9000
    for idx, (label, value, color) in enumerate(rows):
        y = 185 + idx * 82
        draw.text((80, y + 10), label, font=font(19, True), fill="#344054")
        width = (x1 - x0) * value / max_value
        draw.rounded_rectangle((x0, y, x0 + width, y + 48), radius=8, fill=color)
        draw.text((x0 + width + 12, y + 8), f"{value:,}", font=font(19, True), fill="#344054")
    draw.text((80, 690), "신규 개발셋: 96 acoustic profiles · 독립 Test v3: 30 신규 speaker profiles, 교차 중복 0", font=font(18, True), fill="#17365D")
    path = ASSET_DIR / "dataset_before_after.png"
    image.save(path)
    paths["dataset"] = path

    models = EVIDENCE["legacy_study"]["baseline_model_benchmark_validation_120"]
    image, draw = canvas("일반 Whisper Baseline 비교", "동일 Validation 120건 · 제조 용어 Recall", 720)
    x0, x1 = 355, 1485
    for idx, row in enumerate(models):
        y = 190 + idx * 78
        draw.text((80, y + 8), row["model"], font=font(20, True), fill="#344054")
        width = (x1 - x0) * row["recall"]
        color = "#2E74B5" if idx == 0 else "#9CC3E5"
        draw.rounded_rectangle((x0, y, x0 + width, y + 44), radius=7, fill=color)
        draw.text((x0 + width + 12, y + 7), pct(row["recall"]), font=font(18, True), fill="#344054")
    gate_x = x0 + (x1 - x0) * 0.85
    draw.line((gate_x, 165, gate_x, 655), fill="#C00000", width=4)
    draw.text((gate_x - 95, 660), "목표 85%", font=font(18, True), fill="#C00000")
    path = ASSET_DIR / "baseline_recall.png"
    image.save(path)
    paths["baseline"] = path

    image, draw = canvas("LoRA 학습량 및 신규 일반화 Pilot", "Legacy 3·5·10h 곡선과 신규 5h 동일 Validation 비교", 830)
    curves = EVIDENCE["legacy_study"]["lora_learning_curve_validation_600"]
    model_colors = {"medium": "#70AD47", "turbo": "#2E74B5", "large-v3": "#ED7D31"}
    x_positions = {"3h": 250, "5h": 550, "10h": 850}
    for model_name in ("medium", "turbo", "large-v3"):
        rows_for_model = [r for r in curves if r["model"] == model_name]
        points = []
        for row in rows_for_model:
            x = x_positions[row["stage"]]
            y = 620 - (row["recall"] - 0.80) / 0.20 * 390
            points.append((x, y))
            draw.ellipse((x - 8, y - 8, x + 8, y + 8), fill=model_colors[model_name])
            draw.text((x - 38, y - 33), pct(row["recall"]), font=font(15, True), fill=model_colors[model_name])
        draw.line(points, fill=model_colors[model_name], width=5)
    for stage, x in x_positions.items():
        draw.text((x - 18, 650), stage, font=font(18, True), fill="#344054")
    for idx, model_name in enumerate(("medium", "turbo", "large-v3")):
        draw.rectangle((160 + idx * 220, 710, 190 + idx * 220, 730), fill=model_colors[model_name])
        draw.text((200 + idx * 220, 706), model_name, font=font(17, True), fill="#344054")
    pilots = EVIDENCE["generalization_study"]["pilot5h_validation_1200"]
    draw.rounded_rectangle((1030, 190, 1510, 660), radius=14, fill="#F3F7FB", outline="#B8C2CC", width=2)
    draw.text((1070, 225), "신규 일반화 5h", font=font(23, True), fill="#17365D")
    for idx, row in enumerate(pilots):
        y = 320 + idx * 150
        draw.text((1070, y), row["display_name"], font=font(19, True), fill="#344054")
        draw.text((1070, y + 42), f"Recall {pct(row['domain_term_recall'])}", font=font(21, True), fill="#2E74B5")
        draw.text((1070, y + 82), f"CER {pct(row['cer'])} · WER {pct(row['wer'])}", font=font(16), fill="#667085")
    path = ASSET_DIR / "lora_comparison.png"
    image.save(path)
    paths["lora"] = path

    heldout = EVIDENCE["generalization_study"]["heldout_results"][:3]
    image, draw = canvas("Frozen Test 일반화 결과", "동일 고정 설정 · Test v1/v2/v3 각각 1회", 900)
    metrics = [("Recall", "domain_term_recall", 0.85, True), ("CER", "cer", 0.07, False), ("WER", "wer", 0.15, False)]
    for col_idx, (label, key, target, higher) in enumerate(metrics):
        x = 90 + col_idx * 500
        draw.text((x, 180), label, font=font(24, True), fill="#17365D")
        for row_idx, row in enumerate(heldout):
            y = 245 + row_idx * 150
            actual = row[key]
            passed = actual >= target if higher else actual <= target
            color = "#2E7D32" if passed else "#C00000"
            draw.text((x, y), row["cohort"], font=font(18, True), fill="#344054")
            draw.text((x + 150, y - 5), pct(actual), font=font(28, True), fill=color)
            draw.text((x + 310, y + 2), "PASS" if passed else "FAIL", font=font(18, True), fill=color)
            draw.line((x, y + 52, x + 420, y + 52), fill="#E6E9ED", width=2)
        direction = "이상" if higher else "이하"
        draw.text((x, 735), f"목표 {pct(target, 0)} {direction}", font=font(18, True), fill="#667085")
    draw.rounded_rectangle((90, 790, 1510, 860), radius=12, fill="#FDE9E7")
    draw.text((125, 810), "독립 Test v3는 Recall·CER·WER 모두 실패 → 실제 생산 준비 완료 주장을 보류하고 다음 개발 사이클로 전환", font=font(18, True), fill="#C00000")
    path = ASSET_DIR / "heldout_generalization.png"
    image.save(path)
    paths["heldout"] = path

    quant = EVIDENCE["generalization_study"]["quantization_validation_1200"]
    image, draw = canvas("양자화 Trade-off", "Whisper large-v3 LoRA · Validation 1200", 700)
    float_row = next(r for r in quant if r["variant"] == "float16")
    int8_row = next(r for r in quant if r["variant"] == "int8-float16")
    columns = [
        ("모델 크기", gb(float_row["model_size_bytes"]), gb(int8_row["model_size_bytes"]), "49.67% 감소"),
        ("Recall", pct(float_row["domain_term_recall"]), pct(int8_row["domain_term_recall"]), "+0.16%p"),
        ("CER", pct(float_row["cer"]), pct(int8_row["cer"]), "+0.02%p"),
        ("RTF", f"{float_row['rtf']:.3f}", f"{int8_row['rtf']:.3f}", "INT8가 39.3% 느림"),
        ("Peak GPU", f"{float_row['peak_gpu_mb']:.0f} MB", f"{int8_row['peak_gpu_mb']:.0f} MB", "42.1% 감소"),
    ]
    for idx, (label, fp16, int8, delta) in enumerate(columns):
        x = 70 + idx * 300
        draw.rounded_rectangle((x, 190, x + 270, 575), radius=14, fill="#F3F7FB", outline="#B8C2CC", width=2)
        draw_centered(draw, (x + 10, 215, x + 260, 270), label, size=20, bold=True)
        draw.text((x + 30, 320), "FP16", font=font(16, True), fill="#667085")
        draw.text((x + 30, 355), fp16, font=font(23, True), fill="#2E74B5")
        draw.text((x + 30, 425), "INT8-FP16", font=font(16, True), fill="#667085")
        draw.text((x + 30, 460), int8, font=font(23, True), fill="#C55A11")
        draw_centered(draw, (x + 15, 520, x + 255, 565), delta, size=15, fill="#344054", bold=True)
    path = ASSET_DIR / "quantization_tradeoff.png"
    image.save(path)
    paths["quant"] = path
    return paths


def add_picture(doc: Document, path: Path, width=6.6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Inches(width))


def build_report() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    paths = create_charts()
    doc = Document()
    configure_document(doc)

    legacy = EVIDENCE["legacy_study"]
    general = EVIDENCE["generalization_study"]
    targets = EVIDENCE["quality_targets"]
    v1, v2, v3, combined = general["heldout_results"]

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(50)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("AI SPECIALIST  ·  FINAL TECHNICAL REPORT")
    r.font.name = "Aptos"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("제조 현장 한국어 음성인식")
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    r.font.size = Pt(29)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("성능 고도화·일반화 검증·온디바이스 최적화 연구")
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)

    add_table(
        doc,
        ["항목", "내용"],
        [
            ("과제 목표", "제조 용어 Recall≥85%, CER≤7%, WER≤15%의 재현 가능한 한국어 ASR"),
            ("기존 연구", "Whisper Turbo 10h LoRA · 기존 Frozen Test v1 PASS 결과 유지"),
            ("일반화 확장", "Whisper large-v3 pilot-5h LoRA · 신규 Validation 1,200·독립 Test v3"),
            ("검증 환경", "Google Colab Pro · NVIDIA A100-SXM4-80GB · 고용량 RAM · Python 3.12"),
            ("기준 코드", f"Git {EVIDENCE['git_commit']} · codex/whisper-benchmark-quantization"),
            ("작성일", "2026-08-26"),
        ],
        widths=[3.5, 12.7],
    )
    add_callout(
        doc,
        "최종 판단: 파이프라인과 일부 회귀 성능은 입증했으나 독립 일반화 게이트는 미통과",
        f"동결 설정으로 Test v1은 PASS했지만 독립 Test v3는 Recall {pct(v3['domain_term_recall'])}, "
        f"CER {pct(v3['cer'])}, WER {pct(v3['wer'])}로 FAIL했다. 결과를 숨기거나 같은 Test에 재적합하지 않고, "
        "생산 준비 완료 주장을 보류한다. 이 실패는 다음 데이터·모델 개선 사이클의 공식 근거다.",
        kind="warn",
    )
    add_body(doc, "본 보고서는 기존 성과와 신규 일반화 연구를 분리해 기록하며, 모든 수치는 실행 산출물에서 추출했다.", color=GRAY, size=8.8)

    add_section(doc, "Executive Summary")
    add_body(
        doc,
        "본 과제는 일반 한국어 음성인식 모델을 제조 현장 용어·설비 코드·숫자·단위에 맞게 고도화하고, "
        "성능·속도·메모리·배포 제약을 동시에 비교하는 재현 가능한 AIAS 파이프라인을 구축했다. 연구는 "
        "Baseline 비교 → LoRA 학습량 비교 → 디코딩 및 IR/NN 선택 → 양자화 → 동결 Test의 순서로 진행했다.",
    )
    add_table(
        doc,
        ["연구 축", "선택 결과", "핵심 실측", "판단"],
        [
            ("기존 성능 연구", "Whisper Turbo 10h LoRA", "Test v1 Recall 98.13%, CER 0.52%, WER 1.09%", "PASS · 기존 결과 유지"),
            ("신규 일반화 모델", "Whisper large-v3 pilot-5h", "Validation Recall 96.99%, CER 0.24%, WER 0.82%", "Turbo 대비 우세"),
            ("동결 Test v1", "INT8-FP16·no-correction", f"Recall {pct(v1['domain_term_recall'])}, CER {pct(v1['cer'])}, WER {pct(v1['wer'])}", "PASS"),
            ("동결 Test v2", "동일 설정", f"Recall {pct(v2['domain_term_recall'])}, CER {pct(v2['cer'])}, WER {pct(v2['wer'])}", "Recall만 통과"),
            ("독립 Test v3", "동일 설정", f"Recall {pct(v3['domain_term_recall'])}, CER {pct(v3['cer'])}, WER {pct(v3['wer'])}", "FAIL"),
            ("온디바이스 후보", "INT8-FP16", "3.09→1.56GB, GPU 3,572→2,068MB", "용량·메모리 이점, A100 속도 저하"),
        ],
        widths=[3.0, 3.5, 6.1, 3.0],
    )
    add_callout(
        doc,
        "심사 관점의 핵심 가치",
        "높은 숫자만 제시한 것이 아니라, 모델 선정 근거·재현 가능한 CLI·데이터 무결성·Validation/Test 분리·실패 공개·온디바이스 절충을 하나의 연구 체계로 구현했다.",
        kind="info",
    )

    add_section(doc, "1. 추진 배경과 과제 목표")
    doc.add_heading("1.1 현업 문제", level=2)
    add_body(
        doc,
        "제조 현장의 음성 기록은 소음, 짧은 지시문, 유사 발음, 설비 번호, 영문 코드와 단위가 혼재한다. 일반 ASR의 "
        "문장 전체가 자연스러워도 ‘체결→체별’, 설비명·LOT·토크 값이 틀리면 작업 지시·품질 이력에 직접 영향을 준다. "
        "따라서 최우선 KPI를 제조 용어 Recall로 두고 CER·WER를 체감 품질 및 보조 지표로 관리했다.",
    )
    doc.add_heading("1.2 사전 정의한 성공 기준", level=2)
    add_table(
        doc,
        ["지표", "목표", "의미", "선택 규칙"],
        [
            ("제조 용어 Recall", "≥ 85%", "정답 제조 용어를 놓치지 않는 비율", "최우선"),
            ("CER", "≤ 7%", "문자 단위 오류율", "체감 품질"),
            ("WER", "≤ 15%", "어절 단위 오류율", "참고·문장 품질"),
            ("Precision/F1", "함께 보고", "과보정·균형 성능 확인", "Recall 단독 해석 방지"),
            ("RTF/P95/메모리", "장비별 실측", "실시간성·배포 제약", "정확도 통과 후 절충"),
        ],
        widths=[3.2, 2.1, 6.1, 4.2],
    )
    doc.add_heading("1.3 연구 질문", level=2)
    add_bullets(
        doc,
        [
            "어떤 Whisper 크기가 제조 용어 정확도와 추론 효율의 가장 좋은 출발점인가?",
            "LoRA 학습량을 3→5→10시간으로 늘릴 때 성능 곡선과 자원 소요는 어떻게 변하는가?",
            "Hotwords·beam search·VAD와 alias·IR·NN 보정은 실제 Validation에서 이득이 있는가?",
            "FP16과 INT8-FP16 중 정확도 우선 및 온디바이스 후보를 어떻게 분리할 것인가?",
            "기존 Test가 아닌 신규 화자·문장·음향 프로필에서도 성능이 유지되는가?",
        ],
    )

    add_section(doc, "2. 연구 프로토콜과 전체 아키텍처")
    add_picture(doc, paths["flow"])
    add_caption(doc, "그림 1. Validation 선택과 Frozen Test를 분리한 전체 연구 흐름")
    add_body(
        doc,
        "모델·LoRA 학습량·디코딩·후처리·양자화는 Validation에서만 선택한다. 선택 후 configuration을 동결하고 "
        "Test v1/v2/v3를 각각 한 번만 평가했다. Test 결과가 낮더라도 같은 Test에 맞춰 threshold·사전·모델을 다시 "
        "조정하지 않았다. 이는 데이터 누수와 과적합을 방지하는 본 연구의 핵심 통제다.",
    )
    add_picture(doc, paths["architecture"])
    add_caption(doc, "그림 2. 데이터·AIAS CLI·A100·불변 산출물로 연결되는 상세 아키텍처")
    add_callout(
        doc,
        "재현성 원칙",
        "모든 중요한 실행은 config snapshot, environment, manifest, predictions, metrics, run summary, DOCX report와 SQLite registry로 추적한다. Hugging Face 모델은 resolved revision으로 잠근다.",
        kind="info",
    )

    add_section(doc, "3. 데이터 설계와 변경 전후")
    add_picture(doc, paths["dataset"])
    add_caption(doc, "그림 3. 기존 데이터는 보존하고 일반화 개발셋과 독립 Test를 추가한 구조")
    add_table(
        doc,
        ["구분", "Train", "Validation", "Test", "화자/음향", "역할"],
        [
            ("기존 v1 연구", "6,000", "600", "600", "24 speaker profiles", "모델·학습량·최초 frozen 성능"),
            ("일반화 개발", "9,000", "1,200", "-", "96 acoustic profiles", "오류 중심 다양화와 Validation 선택"),
            ("독립 Test v3", "-", "-", "600", "30 신규 profiles", "최종 독립 confirmatory 평가"),
        ],
        widths=[3.0, 2.0, 2.2, 2.0, 3.2, 4.0],
    )
    doc.add_heading("3.1 실패 용어 중심 합성 전략", level=2)
    add_bullets(
        doc,
        [
            "‘체결→체별’과 같이 실제로 관찰된 오인식 쌍을 우선 보강했다.",
            "용어를 여러 문장·발화 속도·음향 profile·숫자·단위 조합에 배치해 단순 반복을 피했다.",
            "영문 코드·설비 번호·단위 표기를 일관되게 정규화했다.",
            "Train/Validation/Test의 문장·speaker profile·audio hash·acoustic profile 중복을 검사했다.",
        ],
    )
    integrity = general["independent_test_v3"]
    add_table(
        doc,
        ["독립 Test v3 무결성 항목", "실측", "판정"],
        [
            ("샘플", f"{integrity['samples']}건", "PASS"),
            ("신규 speaker profiles", f"{integrity['new_speaker_profiles']}개", "PASS"),
            ("제조 용어 출현", f"{integrity['manufacturing_term_occurrences']:,}회", "PASS"),
            ("Negative samples", f"{integrity['negative_samples']}건", "PASS"),
            ("문장·화자·audio hash·acoustic profile 중복", "모두 0", "PASS"),
        ],
        widths=[7.0, 4.0, 4.0],
    )
    add_callout(
        doc,
        "거버넌스 범위",
        "GitHub·Colab에는 AI 생성 합성 음성만 사용했고 개인정보·음성 생체정보·기밀 공장 로그는 반입하지 않았다. 합성 음성은 실제 작업자·소음·마이크 분포를 대체하지 못한다.",
        kind="warn",
    )

    add_section(doc, "4. 관련 Solution 검토와 Baseline 모델 선정", page_break=False)
    add_picture(doc, paths["baseline"])
    add_caption(doc, "그림 4. 일반 Whisper 6종 Baseline의 제조 용어 Recall")
    add_table(
        doc,
        ["모델", "Recall", "CER", "WER", "RTF", "P95", "Peak GPU", "크기"],
        [
            (
                row["model"], pct(row["recall"]), pct(row["cer"]), pct(row["wer"]),
                f"{row['rtf']:.3f}", f"{row['p95_latency_s']:.3f}s", f"{row['peak_gpu_mb']:.0f}MB", gb(row["size_bytes"]),
            )
            for row in legacy["baseline_model_benchmark_validation_120"]
        ],
        widths=[2.2, 2.0, 1.8, 1.8, 1.6, 1.9, 2.1, 2.2],
    )
    add_body(
        doc,
        "일반 모델에서는 large-v3가 Recall 79.86%로 가장 높았지만 목표 85%에는 미달했다. Turbo는 Recall이 "
        "76.98%로 조금 낮은 대신 모델 크기와 RTF가 유리했다. 따라서 large-v3·Turbo·medium을 LoRA 후보로 남겨 "
        "정확도와 자원을 동일 조건에서 재비교했다. tiny/base/small은 경량 후보이나 제조 용어 손실이 커 최종 후보에서 제외했다.",
    )

    add_section(doc, "5. LoRA 파인튜닝과 모델·학습량 비교")
    add_picture(doc, paths["lora"])
    add_caption(doc, "그림 5. 기존 3·5·10시간 학습곡선과 신규 일반화 5시간 Pilot")
    doc.add_heading("5.1 기존 7,200건 연구: 3h→5h→10h", level=2)
    add_table(
        doc,
        ["학습량", "모델", "Train", "학습시간", "Recall", "CER", "WER"],
        [
            (row["stage"], row["model"], f"{row['train_samples']:,}", f"{row['train_seconds']/60:.1f}분", pct(row["recall"]), pct(row["cer"]), pct(row["wer"]))
            for row in legacy["lora_learning_curve_validation_600"]
        ],
        widths=[1.8, 2.5, 2.0, 2.3, 2.2, 2.2, 2.2],
    )
    selected = legacy["selected_lora_paired_validation_600"]
    add_table(
        doc,
        ["동일 Validation 600", "Recall", "CER", "WER"],
        [
            ("Turbo Base", pct(selected["base"]["recall"]), pct(selected["base"]["cer"]), pct(selected["base"]["wer"])),
            ("Turbo 10h LoRA", pct(selected["lora"]["recall"]), pct(selected["lora"]["cer"]), pct(selected["lora"]["wer"])),
            ("개선폭", f"+{(selected['lora']['recall']-selected['base']['recall'])*100:.2f}%p", f"{(selected['lora']['cer']-selected['base']['cer'])*100:.2f}%p", f"{(selected['lora']['wer']-selected['base']['wer'])*100:.2f}%p"),
        ],
        widths=[5.3, 3.4, 3.4, 3.4],
    )
    add_body(
        doc,
        "기존 연구에서는 Turbo 10h를 최종 accuracy-first 모델로 선택했고, 기존 frozen Test v1 성과를 그대로 유지했다. "
        "Whisper Turbo의 공식 ID는 openai/whisper-large-v3-turbo이며 Whisper large-v3와 별개의 후보이다.",
    )
    doc.add_heading("5.2 신규 일반화 연구: 동일 Validation 1,200에서 Pilot-5h", level=2)
    add_table(
        doc,
        ["후보", "Precision", "Recall", "F1", "CER", "WER", "선택"],
        [
            (row["display_name"], pct(row["domain_term_precision"]), pct(row["domain_term_recall"]), pct(row["domain_term_f1"]), pct(row["cer"]), pct(row["wer"]), "선택" if row["display_name"] == "Whisper large-v3" else "-")
            for row in general["pilot5h_validation_1200"]
        ],
        widths=[3.8, 2.0, 2.0, 2.0, 1.8, 1.8, 2.0],
    )
    add_callout(
        doc,
        "신규 일반화 연구의 모델 선택",
        "Whisper large-v3 pilot-5h가 Turbo보다 Recall +1.38%p, CER -0.10%p, WER -0.33%p로 우세하여 선택됐다. 기존 Turbo 10h 결과를 삭제하거나 대체하지 않고 별도 연구 축으로 기록했다.",
        kind="pass",
    )

    add_section(doc, "6. Validation 최적화: 디코딩·IR/NN·양자화", page_break=False)
    doc.add_heading("6.1 디코딩", level=2)
    decode = general["decoding_validation_1200"]
    add_table(
        doc,
        ["선택 설정", "Recall", "CER", "WER", "RTF", "P95", "Peak GPU"],
        [(
            decode["selected"], pct(decode["domain_term_recall"]), pct(decode["cer"]), pct(decode["wer"]),
            f"{decode['rtf']:.3f}", f"{decode['p95_latency_s']:.3f}s", f"{decode['peak_gpu_mb']:.0f}MB",
        )],
        widths=[4.5, 2.0, 1.8, 1.8, 1.8, 2.0, 2.2],
    )
    add_bullets(
        doc,
        [
            "Hotwords: 제조 용어를 디코더에 문맥 힌트로 제공하되 정답을 강제로 덮어쓰지 않는다.",
            "Beam search 8: 후보 문장 8개 경로를 비교해 정확도를 높이는 대신 지연이 증가한다.",
            "VAD 300ms: 짧은 무음 경계로 음성 구간을 분리해 공백·침묵 영향을 줄인다.",
        ],
    )
    doc.add_heading("6.2 Alias·IR·NN 후처리", level=2)
    correction = general["correction_validation_1200"]
    add_table(
        doc,
        ["후보", "변경 건수", "Validation 개선", "결정"],
        [
            ("no-correction", "0", "원문 기준", "선택"),
            ("alias-only", str(correction["alias_only_changed"]), "0", "미적용"),
            ("Information Retrieval", str(correction["ir_changed"]), "0", "미적용"),
            ("Nearest Neighbor", str(correction["nn_changed"]), "0", "미적용"),
            ("hybrid", str(correction["hybrid_changed"]), "0", "미적용"),
        ],
        widths=[4.0, 3.0, 3.5, 4.5],
    )
    add_body(
        doc,
        "IR은 인식문과 용어·문장 사전의 검색 점수를, NN은 임베딩 거리의 근접도를 이용해 후보를 제시한다. 그러나 이번 "
        "Validation에서는 변경·개선 표본이 없어 강제 보정의 근거가 없었다. 따라서 성능을 높이기 위한 기능이라도 검증 이득이 "
        "없으면 적용하지 않는 안전 게이트를 유지했다.",
    )
    doc.add_heading("6.3 양자화", level=2)
    add_picture(doc, paths["quant"])
    add_caption(doc, "그림 6. FP16과 INT8-FP16의 정확도·용량·속도·메모리 비교")
    add_table(
        doc,
        ["Variant", "Recall", "CER", "WER", "RTF", "P95", "크기", "Peak GPU"],
        [
            (row["variant"], pct(row["domain_term_recall"]), pct(row["cer"]), pct(row["wer"]), f"{row['rtf']:.3f}", f"{row['p95_latency_s']:.3f}s", gb(row["model_size_bytes"]), f"{row['peak_gpu_mb']:.0f}MB")
            for row in general["quantization_validation_1200"]
        ],
        widths=[3.0, 2.0, 1.7, 1.7, 1.6, 1.9, 2.2, 2.2],
    )
    add_callout(
        doc,
        "온디바이스 절충",
        "INT8-FP16은 모델 크기를 49.67%, Peak GPU를 42.1% 줄이고 Recall도 +0.16%p였으나, A100 RTF는 0.166→0.231로 악화됐다. 양자화가 항상 더 빠른 것은 아니며 목표 CPU/런타임에서 별도 benchmark가 필요하다.",
        kind="warn",
    )

    add_section(doc, "7. Frozen Test v1·v2·v3 일반화 결과")
    add_picture(doc, paths["heldout"])
    add_caption(doc, "그림 7. 동일 동결 설정의 Test cohort별 품질 게이트")
    add_table(
        doc,
        ["Cohort", "N", "Precision", "Recall", "F1", "CER", "WER", "TP/FP/FN", "Gate"],
        [
            (row["cohort"], row["samples"], pct(row["domain_term_precision"]), pct(row["domain_term_recall"]), pct(row["domain_term_f1"]), pct(row["cer"]), pct(row["wer"]), f"{row['tp']}/{row['fp']}/{row['fn']}", row["quality_gate"])
            for row in general["heldout_results"]
        ],
        widths=[3.2, 1.2, 1.8, 1.8, 1.8, 1.6, 1.6, 2.2, 1.6],
    )
    add_callout(
        doc,
        "독립 Test v3: FAIL",
        f"Recall {pct(v3['domain_term_recall'])} < 85%, CER {pct(v3['cer'])} > 7%, WER {pct(v3['wer'])} > 15%다. "
        "Combined Recall은 87.81%로 목표를 넘지만 CER 13.34%, WER 29.13%가 실패하므로 통합값만으로 성공을 주장하지 않는다.",
        kind="fail",
    )
    doc.add_heading("7.1 Bootstrap 95% 신뢰구간", level=2)
    ci_by_cohort = {}
    for row in general["bootstrap_95_ci"]:
        ci_by_cohort.setdefault(row["cohort"], {})[row["metric"]] = row
    ci_rows = []
    for cohort in ("test_v1", "test_v2", "test_v3", "combined_v1_v2_v3"):
        values = ci_by_cohort[cohort]
        ci_rows.append((
            cohort,
            f"{pct(values['precision']['estimate'])} [{pct(values['precision']['low'])}, {pct(values['precision']['high'])}]",
            f"{pct(values['recall']['estimate'])} [{pct(values['recall']['low'])}, {pct(values['recall']['high'])}]",
            f"{pct(values['f1']['estimate'])} [{pct(values['f1']['low'])}, {pct(values['f1']['high'])}]",
            f"{pct(values['cer']['estimate'])} [{pct(values['cer']['low'])}, {pct(values['cer']['high'])}]",
            f"{pct(values['wer']['estimate'])} [{pct(values['wer']['low'])}, {pct(values['wer']['high'])}]",
        ))
    add_table(doc, ["Cohort", "Precision 95% CI", "Recall 95% CI", "F1 95% CI", "CER 95% CI", "WER 95% CI"], ci_rows, widths=[2.8, 2.8, 2.8, 2.8, 2.8, 2.8])
    add_body(
        doc,
        "독립 Test v3 Recall의 95% CI 상한도 81.29%로 목표 85%에 미치지 못한다. 이는 단순 표본 변동만으로 목표 미달을 "
        "설명하기 어렵다는 근거다. 반대로 v2 Recall은 86.42%이고 CI가 83.76~88.82%로 목표 경계에 걸쳐 있어 현장 판단에는 "
        "추가 표본이 필요하다.",
    )
    doc.add_heading("7.2 Historical v2 paired 비교", level=2)
    add_table(
        doc,
        ["지표", "이전 Turbo 10h", "신규 large-v3 5h", "변화", "해석"],
        [
            ("Precision", "97.90%", "92.90%", "-5.00%p", "과검출 증가"),
            ("Recall", "82.33%", "86.42%", "+4.09%p", "목표 85% 초과"),
            ("F1", "89.44%", "89.54%", "+0.10%p", "균형 성능 유지"),
            ("CER", "17.35%", "10.56%", "-6.79%p", "개선, 목표는 미달"),
            ("WER", "36.64%", "24.86%", "-11.78%p", "개선, 목표는 미달"),
        ],
        widths=[2.5, 3.0, 3.2, 2.5, 4.8],
    )
    add_body(
        doc,
        "신규 데이터·large-v3 LoRA는 v2의 Recall·CER·WER를 의미 있게 개선했지만 Precision이 하락했고, CER/WER는 여전히 "
        "현업 목표를 통과하지 못했다. 즉 ‘일반화 회복 효과가 있다’는 결론은 가능하지만 ‘현업 적용 가능 수준을 달성했다’는 "
        "결론은 불가능하다.",
    )

    add_section(doc, "8. 실패 분석과 다음 개선 사이클")
    add_table(
        doc,
        ["관찰", "근거", "가능 원인", "다음 검증"],
        [
            ("v1 PASS, v3 FAIL", "동일 동결 설정에서 cohort 차이", "음향·문장 분포 이동", "v3 오류 군집을 development에만 편입"),
            ("Recall보다 CER/WER 악화가 큼", "v2/v3 문장 오류율 상승", "문장 구조·숫자·코드 decoding 취약", "숫자·단위·설비 코드별 stratified eval"),
            ("INT8가 A100에서 느림", "RTF 0.166→0.231", "GPU kernel·메모리 전송 절충", "목표 CPU 및 RTX 6000에서 재측정"),
            ("IR/NN 무효", "Validation changed/improved 0", "후보 사전·threshold 범위 부족 또는 불필요", "오류 로그로 후보군 재설계"),
            ("합성→실환경 간극 미측정", "실제 작업자 데이터 없음", "마이크·소음·억양 차이", "승인된 real-world shadow cohort"),
        ],
        widths=[3.2, 3.8, 4.2, 4.8],
    )
    doc.add_heading("8.1 Test를 오염시키지 않는 개선 방법", level=2)
    add_bullets(
        doc,
        [
            "이번 v3를 반복 튜닝용 Validation으로 전환하지 않는다. 결과는 고정 증빙으로 보존한다.",
            "v3의 오류 유형을 집계해 별도의 신규 development 데이터·Validation 후보를 만든다.",
            "모델·디코딩·보정 선택 후 새로운 신규 화자·문장·음성의 Test v4를 사전 등록하고 1회 평가한다.",
            "Recall·CER·WER뿐 아니라 Precision/F1, 숫자·단위·설비 코드 slice, 지연·메모리까지 함께 gate한다.",
        ],
    )
    doc.add_heading("8.2 연구 마무리 시점의 합리적 결론", level=2)
    add_callout(
        doc,
        "현재 확보한 성과",
        "재현 가능한 전체 파이프라인, 일반 모델/LoRA/디코딩/후처리/양자화 비교, 기존 v1의 높은 성능, v2 회복 효과, 독립 v3의 정직한 실패 및 신뢰구간까지 확보했다.",
        kind="pass",
    )
    add_callout(
        doc,
        "현재 확보하지 못한 성과",
        "신규 화자·음향 분포에서 세 품질 게이트를 동시에 통과하는 일반화 성능과 실제 공장 production readiness는 아직 입증하지 못했다.",
        kind="fail",
    )

    add_section(doc, "9. AI Application 및 On-device 배포 검토")
    doc.add_heading("9.1 RTX 6000 Ada 2장 Workstation", level=2)
    add_body(
        doc,
        "Whisper는 autoregressive ASR이며 일반 LLM 전용 vLLM 구조에 그대로 올리는 것이 최적 경로는 아니다. RTX 6000 Ada "
        "48GB×2에서는 faster-whisper/CTranslate2 기반의 독립 worker를 GPU별로 배치하고, API router가 요청을 분산하는 구조가 "
        "운영·성능 측면에서 단순하다. 한 요청을 두 GPU에 분할하기보다 모델 복제와 동시성 확장이 적합하다.",
    )
    add_table(
        doc,
        ["항목", "RTX 6000 Ada ×2", "권장 판단"],
        [
            ("정확도 우선", "FP16 large-v3 LoRA", "GPU 메모리 여유가 커 품질 우선 서빙"),
            ("동시성", "GPU별 worker + queue/router", "배치·동시 사용자 수 실측"),
            ("양자화", "INT8-FP16 선택 가능", "A100에서 느렸으므로 Ada에서 재검증"),
            ("관측", "RTF, P95, GPU memory, queue latency", "부하 수준별 SLO 설정"),
        ],
        widths=[3.3, 6.2, 6.3],
    )
    doc.add_heading("9.2 GPU 없는 CPU On-device", level=2)
    add_body(
        doc,
        "CPU 단독 장비에서는 large-v3 INT8의 용량 1.56GB 자체는 탑재 가능해도 실시간성과 전력·발열이 병목이 될 수 있다. "
        "AVX2/AVX-512 지원 CPU, 물리 코어 수, 메모리 대역폭, 긴 음성 길이에 따라 편차가 크다. 따라서 medium/small INT8을 "
        "별도 후보로 두고 목표 장비에서 RTF<1, P95, RSS, 전력, thermal throttling을 측정해야 한다.",
    )
    add_table(
        doc,
        ["비교", "RTX 6000 Ada", "CPU On-device"],
        [
            ("정확도 후보", "large-v3 FP16/INT8", "medium/small INT8부터 검토"),
            ("성능", "높은 처리량·동시성", "코어·SIMD·메모리에 민감"),
            ("개인정보", "중앙 서버 전송 필요", "로컬 처리로 유리"),
            ("운영", "서버·큐·모니터링", "장비별 업데이트·열·전력 관리"),
            ("현재 근거", "A100 proxy benchmark", "실장비 미측정"),
        ],
        widths=[3.3, 6.2, 6.3],
    )

    add_section(doc, "10. 프로그래밍·Tool 활용과 재현성")
    add_table(
        doc,
        ["기능", "구현/Tool", "산출물"],
        [
            ("환경", "Python 3.12 · uv · pinned dependencies", "environment.json"),
            ("데이터", "manifest parser · ffmpeg · hash/integrity checks", "prepared_manifest.csv"),
            ("모델", "Hugging Face model lock · Transformers/PEFT LoRA", "model-lock.yaml · adapters"),
            ("Benchmark/Optimize", "faster-whisper · CTranslate2 · FP16/INT8", "comparison.csv"),
            ("후처리", "alias · Information Retrieval · Nearest Neighbor", "correction_audit.csv"),
            ("검증", "immutable run · config snapshot · SQLite registry", "metrics.json · run_summary.md"),
            ("보고", "bootstrap CI · DOCX 자동 생성 · render QA", "최종 기술 보고서"),
        ],
        widths=[3.2, 6.3, 6.3],
    )
    add_body(
        doc,
        "AIAS CLI는 GitHub 저장소의 aias_specialist.cli 모듈을 Python -m으로 실행하는 인터페이스다. prepare, benchmark-models, "
        "train-selected-whisper, decoding sweep, correction sweep, quantization, final evaluation 같은 명령을 공통 함수로 호출해 "
        "Colab 셀의 중복 코드를 줄이고, 동일한 로그·오류 처리·환경을 재사용한다.",
    )
    add_callout(
        doc,
        "불변 실행(immutable run)",
        "한 번 완료된 실행 폴더와 config snapshot을 덮어쓰지 않고 새 run_id로 생성한다. 따라서 어떤 코드·모델·설정·데이터로 나온 수치인지 역추적할 수 있다.",
        kind="info",
    )

    add_section(doc, "11. AI Governance와 사람 검토 Gate")
    add_table(
        doc,
        ["위험", "현재 통제", "생산 전 추가 조치"],
        [
            ("음성 개인정보·생체정보", "합성 음성만 GitHub/Colab 사용", "동의·목적·보관기간·철회 절차"),
            ("기밀 제조정보", "실제 공장 로그 미반입", "승인된 비식별·최소수집"),
            ("라벨 오류", "manifest·정규화·무결성 검사", "현업 2인 검토 및 sign-off"),
            ("생성 합성 편향", "화자·음향 profile 분리", "실제 shadow cohort로 외적 타당성 검증"),
            ("과보정", "Validation 개선/악화 gate", "confidence·distance·human review"),
            ("과장된 결론", "독립 v3 실패·CI 공개", "생산 준비 완료 표현 금지"),
        ],
        widths=[4.0, 5.8, 6.0],
    )
    add_callout(
        doc,
        "Human review gate",
        "실제 적용 전 transcript label, 제조 용어 사전, 개인정보 승인, 모델·양자화 trade-off, 실패 cohort와 최종 보고서 결론을 현업 책임자가 검토·서명해야 한다.",
        kind="warn",
    )

    add_section(doc, "12. 심사기준별 성과와 최종 결론")
    add_table(
        doc,
        ["심사기준", "구현·연구 근거", "자체 판단"],
        [
            ("1. AI기본·AI모델링", "6개 Baseline, 3개 LoRA 후보, 학습량 곡선, Validation 선택, 독립 Test", "모델 선정 근거와 실패 분석 확보"),
            ("2. 프로그래밍·Tool", "Python 3.12, uv, AIAS CLI, HF, PEFT, CTranslate2, ffmpeg, pytest/ruff", "재현 가능한 자동화 구현"),
            ("3. AI Application·On-device", "A100 전체 실행, 디코딩/IR/NN/양자화, RTX6000/CPU 배포 설계", "시스템 제약을 수치로 비교"),
            ("4. AI Governance", "합성 데이터, provenance, split/hash 검사, immutable run, human gate", "합성 한계를 명시하고 생산 주장 보류"),
        ],
        widths=[4.0, 8.0, 4.0],
    )
    doc.add_heading("12.1 최종 연구 성과", level=2)
    add_bullets(
        doc,
        [
            "Baseline부터 보고서까지 자동화된 제조 한국어 ASR 연구 파이프라인을 구축했다.",
            "기존 Turbo 10h 연구에서 frozen Test v1 품질 게이트 PASS를 확보·보존했다.",
            "신규 일반화 연구에서 large-v3 5h가 Turbo 5h보다 우수함을 Validation으로 선택했다.",
            "v2에서 Recall +4.09%p, CER -6.79%p, WER -11.78%p의 회복 효과를 확인했다.",
            "독립 v3 실패와 95% CI를 공개해 일반화 한계를 수치로 규명했다.",
            "INT8-FP16으로 크기 49.67%, GPU 메모리 42.1% 절감 가능성을 확인했다.",
        ],
    )
    doc.add_heading("12.2 결론", level=2)
    add_callout(
        doc,
        "결론",
        "본 연구는 단순 데모를 넘어 모델 비교·LoRA·디코딩·검색 보정·양자화·통계 검증·거버넌스가 연결된 기술 체계를 확보했다. 그러나 독립 Test v3의 세 품질 게이트 실패로 합성 분포 일반화가 충분하지 않으며, 실제 공장 production readiness는 주장하지 않는다. 다음 단계는 오류 중심 development 재설계와 새로운 사전 등록 Test v4, 승인된 실음성 shadow 평가다.",
        kind="warn",
    )

    add_section(doc, "Appendix A. 핵심 기술 용어 요약")
    glossary = [
        ("Hotwords", "제조 용어가 디코딩 후보에 더 잘 나타나도록 주는 문맥 힌트. 정답 강제 치환은 아님."),
        ("Beam search", "여러 토큰 경로를 동시에 비교해 더 좋은 문장을 선택하는 탐색. beam이 커지면 정확도와 지연이 함께 증가 가능."),
        ("VAD", "음성/무음을 구분해 처리할 구간을 나누는 Voice Activity Detection."),
        ("IR", "인식문과 사전 문장·용어의 검색 점수를 이용해 보정 후보를 찾는 Information Retrieval."),
        ("NN Search", "임베딩 공간에서 가장 가까운 후보를 찾는 Nearest Neighbor 검색."),
        ("Held-out Test", "학습·선택에 쓰지 않고 설정 동결 뒤 한 번만 보는 최종 평가셋."),
        ("Alias", "같은 제조 개념의 표기·발음 변형을 표준 용어로 연결하는 사전."),
        ("Manifest", "음성 경로, 정답, split, 화자, 출처, 승인 상태를 행 단위로 기록한 데이터 목록."),
        ("ffmpeg", "음성 형식·sample rate·channel 변환 및 길이 확인에 쓰는 미디어 도구."),
        ("CTranslate2", "Transformer 모델을 CPU/GPU에서 효율적으로 추론하고 FP16/INT8로 변환하는 런타임."),
        ("Immutable run", "기존 실행을 덮어쓰지 않고 새 run_id와 산출물 폴더로 보존하는 실행 단위."),
        ("Config snapshot", "실행 시점의 설정을 복사해 어떤 조건으로 나온 결과인지 박제한 파일."),
        ("RTF", "처리시간÷음성길이. 1보다 작으면 평균적으로 실시간보다 빠름."),
        ("P95 latency", "요청의 95%가 이 시간 안에 끝나는 지연 상한 지표."),
        ("Real-world shadow 평가", "실제 운영 흐름에서 결과를 업무에 반영하지 않고 모델 성능만 관찰하는 안전 평가."),
        ("현장 cohort 평가", "공정·소음·마이크·화자 등 의미 있는 현장 집단별로 성능을 분리해 보는 평가."),
        ("Paired 비교", "동일 sample에 두 모델/설정을 적용해 샘플 단위 변화가 우연인지 비교하는 방법."),
    ]
    add_table(doc, ["용어", "요약"], glossary, widths=[4.0, 12.0])

    add_section(doc, "Appendix B. 실행 증빙과 추적 경로")
    sources = general["source_paths"]
    add_table(
        doc,
        ["증빙", "경로/식별자"],
        [
            ("Git 기준", f"{EVIDENCE['git_commit']} / codex/whisper-benchmark-quantization"),
            ("기존 Test v1 run", legacy["frozen_test_v1_original"]["run_id"]),
            ("일반화 비교", sources["comparison"]),
            ("Bootstrap CI", sources["confidence_intervals"]),
            ("v2 before/after", sources["v2_before_after"]),
            ("양자화 비교", sources["quantization"]),
            ("로컬 근거 snapshot", str(EVIDENCE_PATH.relative_to(ROOT)).replace("\\", "/")),
        ],
        widths=[4.0, 12.0],
    )
    doc.add_heading("B.1 수치 해석 규칙", level=2)
    add_bullets(doc, EVIDENCE["protocol_notes"])
    add_body(
        doc,
        "Precision=TP/(TP+FP), Recall=TP/(TP+FN), F1=2PR/(P+R). Combined는 cohort별 백분율의 단순 평균이 아니라 "
        "TP·FP·FN을 합산한 micro aggregate다. CER/WER는 동일 정규화 규칙으로 계산했다.",
    )

    core = doc.core_properties
    core.title = "AI Specialist 제조 현장 한국어 ASR 최종 기술 보고서"
    core.subject = "성능 고도화, 일반화 검증, 온디바이스 최적화"
    core.author = "AI Specialist ASR Project"
    core.keywords = "ASR, Whisper, Korean, manufacturing, LoRA, generalization, quantization, governance"
    core.comments = "Evidence snapshot 2026-08-26; synthetic-data study"

    doc.save(DOCX_PATH)
    print(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    build_report()
