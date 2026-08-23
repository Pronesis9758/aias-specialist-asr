from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from docx import Document

from .config import load_settings
from .correction import apply_term_correction
from .correction_audit import write_correction_audit
from .evaluation import compare_metrics, evaluate_predictions, per_sample_metrics
from .quality_targets import quality_target_columns
from .utils import new_run_id, utc_now


@dataclass(frozen=True)
class CorrectionSweepResult:
    sweep_id: str
    sweep_dir: Path
    comparison_path: Path
    report_path: Path
    recommended_candidate: str


def _load_yaml(path: Path) -> dict[str, Any]:
    payload = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"YAML must contain a mapping: {path}")
    return payload


def _safe_id(value: Any, label: str) -> str:
    identifier = str(value or "").strip()
    if not identifier or any(character in identifier for character in '\\/:*?"<>|'):
        raise ValueError(f"{label} must be a non-empty filesystem-safe identifier")
    return identifier


def _source_run(selection_path: Path) -> tuple[dict[str, Any], Path, pd.Series]:
    selection = _load_yaml(selection_path).get("selection")
    if not isinstance(selection, dict) or selection.get("source_kind") != "benchmark":
        raise ValueError("Correction sweep requires a model_selection.yaml from a benchmark")
    comparison_path = selection_path.parent / "benchmark_comparison.csv"
    comparison = pd.read_csv(comparison_path, keep_default_na=False)
    matches = comparison.loc[comparison["run_id"].astype(str) == str(selection["run_id"])]
    if len(matches) != 1:
        raise ValueError("Selected benchmark run is missing or ambiguous")
    row = matches.iloc[0]
    if str(row["evaluation_split"]).lower() != "validation":
        raise ValueError("Correction tuning must use validation predictions, never held-out test")
    run_dir = Path(str(row["run_dir"])).expanduser().resolve()
    return selection, run_dir, row


def _candidate_options(candidate: dict[str, Any]) -> dict[str, Any]:
    information_retrieval = candidate.get("information_retrieval", {})
    nearest_neighbor = candidate.get("nearest_neighbor", {})
    if not isinstance(information_retrieval, dict) or not isinstance(nearest_neighbor, dict):
        raise ValueError("Correction candidate retrieval settings must be mappings")
    return {
        "enabled": bool(candidate.get("enabled", True)),
        "case_sensitive": bool(candidate.get("case_sensitive", False)),
        "alias_enabled": bool(candidate.get("alias_enabled", True)),
        "information_retrieval_enabled": bool(information_retrieval.get("enabled", False)),
        "nearest_neighbor_enabled": bool(nearest_neighbor.get("enabled", False)),
        "top_k": int(candidate.get("top_k", 1)),
        "max_ngram_tokens": int(candidate.get("max_ngram_tokens", 2)),
        "min_ir_score": float(information_retrieval.get("min_score", 0.8)),
        "min_nn_score": float(nearest_neighbor.get("min_score", 0.9)),
        "nn_backend": str(nearest_neighbor.get("backend", "char_ngram")),
        "nn_model_repo_id": nearest_neighbor.get("model_repo_id"),
        "nn_model_revision": str(nearest_neighbor.get("model_revision", "main")),
        "nn_device": str(nearest_neighbor.get("device", "auto")),
        "ir_weight": float(information_retrieval.get("weight", 0.5)),
        "nn_weight": float(nearest_neighbor.get("weight", 0.5)),
        "require_consensus": bool(candidate.get("require_consensus", False)),
        "min_score_margin": float(candidate.get("min_score_margin", 0.0)),
        "max_length_ratio": float(candidate.get("max_length_ratio", 4.0)),
    }


def _acceptance(
    row: dict[str, Any],
    gates: dict[str, Any],
    *,
    fallback: bool,
) -> tuple[bool, str]:
    if fallback:
        return True, "safe baseline fallback"
    checks = {
        "CER reduction": row["cer_absolute_reduction"]
        >= float(gates.get("min_cer_absolute_reduction", 1e-9)),
        "WER reduction": row["wer_absolute_reduction"]
        >= float(gates.get("min_wer_absolute_reduction", 0.0)),
        "term recall gain": row["domain_term_recall_gain"]
        >= float(gates.get("min_domain_term_recall_gain", 0.0)),
        "degraded sample rate": row["degraded_sample_rate"]
        <= float(gates.get("max_degraded_sample_rate", 0.0)),
    }
    failures = [name for name, passed in checks.items() if not passed]
    return not failures, "passed all gates" if not failures else "failed: " + ", ".join(failures)


def _write_report(
    comparison: pd.DataFrame,
    output_path: Path,
    *,
    title: str,
    sweep_id: str,
    source_run_id: str,
    recommended_candidate: str,
) -> Path:
    document = Document()
    document.add_heading(title, level=0)
    document.add_paragraph(f"Correction sweep: {sweep_id}")
    document.add_paragraph(f"Validation source run: {source_run_id}")
    document.add_paragraph(f"Recommended candidate: {recommended_candidate}")
    document.add_heading("1. Validation comparison", level=1)
    table = document.add_table(rows=1, cols=9)
    table.style = "Table Grid"
    headers = [
        "Rank",
        "Candidate",
        "Accepted",
        "CER before",
        "CER after",
        "CER reduction",
        "WER reduction",
        "Term recall gain",
        "Degraded rate",
    ]
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = header
    for row in comparison.itertuples(index=False):
        cells = table.add_row().cells
        values = [
            "" if pd.isna(row.rank) else str(int(row.rank)),
            str(row.candidate_id),
            str(bool(row.accepted)),
            f"{float(row.cer_before):.4f}",
            f"{float(row.cer_after):.4f}",
            f"{float(row.cer_absolute_reduction):+.4f}",
            f"{float(row.wer_absolute_reduction):+.4f}",
            f"{float(row.domain_term_recall_gain):+.4f}",
            f"{float(row.degraded_sample_rate):.1%}",
        ]
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    document.add_heading("2. Acceptance policy", level=1)
    document.add_paragraph(
        "Only validation configurations that improve CER, do not regress WER or domain-term "
        "recall, and satisfy the degraded-sample guardrail may be selected. If none pass, the "
        "no-correction candidate is selected automatically. Held-out Test data is not used for "
        "threshold tuning."
    )
    document.add_heading("3. Human review gate", level=1)
    document.add_paragraph(
        "Every changed or degraded transcript in the candidate audit files requires review. "
        "Synthetic or public proxy results are not production-readiness evidence."
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = title
    document.core_properties.subject = "ASR correction validation sweep"
    document.save(output_path)
    return output_path


def run_correction_sweep(
    spec_path: str | Path,
    selection_path: str | Path,
) -> CorrectionSweepResult:
    spec_path = Path(spec_path).expanduser().resolve()
    selection_path = Path(selection_path).expanduser().resolve()
    spec = _load_yaml(spec_path)
    section = spec.get("correction_sweep")
    if not isinstance(section, dict):
        raise ValueError("Correction sweep spec requires a correction_sweep mapping")
    if str(section.get("evaluation_split", "validation")).lower() != "validation":
        raise ValueError("Correction sweep must use validation data")
    sweep_id = _safe_id(section.get("id"), "correction_sweep.id")
    base_config = Path(str(section["base_config"])).expanduser()
    if not base_config.is_absolute():
        base_config = (spec_path.parent.parent.parent / base_config).resolve()
    settings = load_settings(base_config)
    selection, source_run_dir, source_row = _source_run(selection_path)
    predictions_path = source_run_dir / "predictions_baseline.csv"
    terms_path = source_run_dir / "domain_terms.snapshot.csv"
    if not predictions_path.exists() or not terms_path.exists():
        raise ValueError("Selected run is missing baseline predictions or domain-term snapshot")
    predictions = pd.read_csv(predictions_path, keep_default_na=False)
    predictions = per_sample_metrics(predictions)
    terms = pd.read_csv(terms_path, keep_default_na=False)
    baseline_metrics = evaluate_predictions(predictions, terms)

    sweep_dir = settings.paths.artifacts_dir.parent / "correction" / sweep_id
    sweep_dir.mkdir(parents=True, exist_ok=True)
    snapshot_path = sweep_dir / "correction_sweep.snapshot.yaml"
    snapshot = {
        **spec,
        "source": {
            "run_id": str(selection["run_id"]),
            "model_id": str(selection["model_id"]),
            "repo_id": str(selection["model"]["repo_id"]),
            "model_revision": str(source_row["model_revision"]),
        },
    }
    serialized_snapshot = yaml.safe_dump(snapshot, allow_unicode=True, sort_keys=False)
    if snapshot_path.exists() and snapshot_path.read_text(encoding="utf-8") != serialized_snapshot:
        raise ValueError(
            f"Correction sweep ID already exists with a different specification: {sweep_id}"
        )
    snapshot_path.write_text(serialized_snapshot, encoding="utf-8")

    raw_candidates = section.get("candidates", [])
    if not isinstance(raw_candidates, list) or not raw_candidates:
        raise ValueError("Correction sweep requires at least one candidate")
    gates = section.get("acceptance", {})
    if not isinstance(gates, dict):
        raise ValueError("correction_sweep.acceptance must be a mapping")
    rows: list[dict[str, Any]] = []
    for raw_candidate in raw_candidates:
        if not isinstance(raw_candidate, dict):
            raise ValueError("Each correction candidate must be a mapping")
        candidate_id = _safe_id(raw_candidate.get("id"), "candidate.id")
        candidate_dir = sweep_dir / "candidates" / candidate_id
        candidate_dir.mkdir(parents=True, exist_ok=True)
        options = _candidate_options(raw_candidate)
        corrected = apply_term_correction(predictions, terms, **options)
        corrected = per_sample_metrics(corrected)
        corrected.to_csv(
            candidate_dir / "predictions_corrected.csv", index=False, encoding="utf-8-sig"
        )
        audit = write_correction_audit(candidate_dir, predictions, corrected)
        corrected_metrics = evaluate_predictions(corrected, terms)
        improvement = compare_metrics(baseline_metrics, corrected_metrics)
        degraded_count = int((audit["outcome"] == "degraded").sum())
        fallback = bool(raw_candidate.get("fallback", False))
        row = {
            "candidate_id": candidate_id,
            "fallback": fallback,
            "sample_count": len(audit),
            "changed_count": int(audit["text_changed"].sum()),
            "improved_count": int((audit["outcome"] == "improved").sum()),
            "degraded_count": degraded_count,
            "degraded_sample_rate": degraded_count / max(len(audit), 1),
            "wer_before": float(baseline_metrics["wer"]),
            "wer_after": float(corrected_metrics["wer"]),
            "wer_absolute_reduction": float(improvement["wer_absolute_reduction"]),
            "cer_before": float(baseline_metrics["cer"]),
            "cer_after": float(corrected_metrics["cer"]),
            "cer_absolute_reduction": float(improvement["cer_absolute_reduction"]),
            "domain_term_recall_before": float(baseline_metrics["domain_term_recall"]),
            "domain_term_recall_after": float(corrected_metrics["domain_term_recall"]),
            "domain_term_recall_gain": float(improvement["domain_term_recall_gain"]),
            "config_path": str(candidate_dir / "correction.yaml"),
        }
        row.update(quality_target_columns(corrected_metrics, settings.quality_targets))
        accepted, reason = _acceptance(row, gates, fallback=fallback)
        row["accepted"] = accepted
        row["acceptance_reason"] = reason
        (candidate_dir / "correction.yaml").write_text(
            yaml.safe_dump({"correction": raw_candidate}, allow_unicode=True, sort_keys=False),
            encoding="utf-8",
        )
        rows.append(row)

    comparison = pd.DataFrame(rows)
    eligible = comparison.loc[comparison["accepted"] & ~comparison["fallback"]].copy()
    if eligible.empty:
        eligible = comparison.loc[comparison["fallback"]].copy()
    if eligible.empty:
        raise ValueError("No accepted correction candidate and no safe fallback candidate")
    if settings.quality_targets.enabled:
        eligible = eligible.sort_values(
            [
                "quality_target_pass",
                "domain_term_recall_gap",
                "cer_gap",
                "wer_gap",
                "degraded_sample_rate",
                "changed_count",
                "candidate_id",
            ],
            ascending=[False, True, True, True, True, False, True],
        )
    else:
        eligible = eligible.sort_values(
            ["cer_after", "wer_after", "degraded_sample_rate", "changed_count", "candidate_id"]
        )
    recommended_candidate = str(eligible.iloc[0]["candidate_id"])
    comparison["rank"] = pd.NA
    ranking = pd.concat(
        [
            eligible,
            comparison.loc[~comparison.index.isin(eligible.index)].sort_values(
                (
                    ["accepted", "domain_term_recall_gap", "cer_gap", "wer_gap"]
                    if settings.quality_targets.enabled
                    else ["accepted", "cer_after", "wer_after"]
                ),
                ascending=(
                    [False, True, True, True]
                    if settings.quality_targets.enabled
                    else [False, True, True]
                ),
            ),
        ]
    )
    for rank, index in enumerate(ranking.index, start=1):
        comparison.loc[index, "rank"] = rank
    comparison = comparison.sort_values("rank")
    comparison_path = sweep_dir / "correction_sweep_comparison.csv"
    comparison.to_csv(comparison_path, index=False, encoding="utf-8-sig")
    report_path = _write_report(
        comparison,
        sweep_dir / "reports/correction_sweep_report.docx",
        title=str(section.get("name", "ASR correction validation sweep")),
        sweep_id=sweep_id,
        source_run_id=str(selection["run_id"]),
        recommended_candidate=recommended_candidate,
    )
    return CorrectionSweepResult(
        sweep_id=sweep_id,
        sweep_dir=sweep_dir,
        comparison_path=comparison_path,
        report_path=report_path,
        recommended_candidate=recommended_candidate,
    )


def select_correction_candidate(
    sweep_dir: str | Path,
    candidate_id: str | None,
    *,
    reviewer: str,
    reason: str,
    human_reviewed: bool,
) -> Path:
    directory = Path(sweep_dir).expanduser().resolve()
    comparison = pd.read_csv(directory / "correction_sweep_comparison.csv", keep_default_na=False)
    if candidate_id is None:
        accepted = comparison.loc[comparison["accepted"].astype(str).str.lower().eq("true")]
        if accepted.empty:
            raise ValueError("Correction sweep has no accepted candidate")
        candidate_id = str(accepted.sort_values("rank").iloc[0]["candidate_id"])
    matches = comparison.loc[comparison["candidate_id"].astype(str) == candidate_id]
    if len(matches) != 1:
        raise ValueError(f"Unknown correction candidate: {candidate_id}")
    row = matches.iloc[0]
    if str(row["accepted"]).lower() != "true":
        raise ValueError(f"Correction candidate did not pass acceptance gates: {candidate_id}")
    candidate_config = _load_yaml(Path(str(row["config_path"])))
    selection_id = f"correction-selection-{new_run_id()}"
    payload = {
        "selection": {
            "selection_id": selection_id,
            "source_kind": "correction_sweep",
            "source_group_id": directory.name,
            "candidate_id": candidate_id,
            "selected_at": utc_now().isoformat(),
            "reviewer": reviewer,
            "reason": reason,
            "human_reviewed": human_reviewed,
            "selection_scope": "human_review" if human_reviewed else "automated_public_proxy",
            "correction": candidate_config["correction"],
            "metrics": {
                key: row[key].item() if hasattr(row[key], "item") else row[key]
                for key in [
                    "wer_before",
                    "wer_after",
                    "wer_absolute_reduction",
                    "cer_before",
                    "cer_after",
                    "cer_absolute_reduction",
                    "domain_term_recall_gain",
                    "degraded_sample_rate",
                ]
            },
            "final_test_required": True,
        }
    }
    serialized = yaml.safe_dump(payload, allow_unicode=True, sort_keys=False)
    immutable = directory / "selections" / f"{selection_id}.yaml"
    immutable.parent.mkdir(parents=True, exist_ok=True)
    immutable.write_text(serialized, encoding="utf-8")
    output_path = directory / "correction_selection.yaml"
    output_path.write_text(serialized, encoding="utf-8")
    return output_path
