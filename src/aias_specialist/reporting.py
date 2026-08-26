from __future__ import annotations

import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT_GRAY_HEX = "F2F4F7"
CALLOUT_HEX = "F4F6F9"


def create_metrics_chart(metrics: dict[str, Any], output_path: Path) -> Path:
    matplotlib_cache = Path(tempfile.gettempdir()) / "aias-matplotlib"
    matplotlib_cache.mkdir(parents=True, exist_ok=True)
    os.environ.setdefault("MPLCONFIGDIR", str(matplotlib_cache))
    from matplotlib.figure import Figure

    output_path.parent.mkdir(parents=True, exist_ok=True)
    if "lora" in metrics:
        return _create_adaptation_metrics_chart(
            metrics, output_path, Figure, "lora", "Best LoRA", "LoRA"
        )
    if "distilled" in metrics:
        return _create_adaptation_metrics_chart(
            metrics, output_path, Figure, "distilled", "Distilled Student", "Distillation"
        )
    labels = ["WER", "CER"]
    baseline = [
        metrics["baseline"]["wer"],
        metrics["baseline"]["cer"],
    ]
    corrected = [
        metrics["corrected"]["wer"],
        metrics["corrected"]["cer"],
    ]
    if metrics["baseline"].get("domain_term_recall_applicable"):
        labels.append("Term Recall")
        baseline.append(metrics["baseline"]["domain_term_recall"])
        corrected.append(metrics["corrected"]["domain_term_recall"])
    positions = list(range(len(labels)))
    width = 0.34

    fig = Figure(figsize=(8, 4.2))
    axis = fig.subplots()
    axis.bar([item - width / 2 for item in positions], baseline, width, label="Baseline")
    axis.bar([item + width / 2 for item in positions], corrected, width, label="Corrected")
    axis.set_xticks(positions, labels)
    axis.set_ylim(0, max(1.0, *(baseline + corrected)) * 1.08)
    axis.set_ylabel("Score")
    axis.set_title("ASR Evaluation: Baseline vs. Domain-term Correction")
    axis.grid(axis="y", alpha=0.25)
    axis.legend()
    fig.tight_layout()
    fig.savefig(output_path, dpi=180, bbox_inches="tight")
    return output_path


def _create_adaptation_metrics_chart(
    metrics: dict[str, Any],
    output_path: Path,
    figure_class: Any,
    variant_key: str,
    variant_label: str,
    method_label: str,
) -> Path:
    baseline = metrics["baseline"]
    adapted = metrics[variant_key]
    fig = figure_class(figsize=(9, 4.2))
    accuracy_axis, speed_axis = fig.subplots(1, 2)

    labels = ["WER", "CER"]
    positions = [0, 1]
    width = 0.34
    baseline_bars = accuracy_axis.bar(
        [position - width / 2 for position in positions],
        [baseline["wer"], baseline["cer"]],
        width,
        label="Base Whisper",
    )
    adapted_bars = accuracy_axis.bar(
        [position + width / 2 for position in positions],
        [adapted["wer"], adapted["cer"]],
        width,
        label=variant_label,
    )
    accuracy_axis.set_xticks(positions, labels)
    accuracy_axis.set_ylim(
        0,
        max(1.0, baseline["wer"], baseline["cer"], adapted["wer"], adapted["cer"]) * 1.16,
    )
    accuracy_axis.set_ylabel("Error rate (lower is better)")
    accuracy_axis.set_title("Accuracy on identical test samples")
    accuracy_axis.grid(axis="y", alpha=0.25)
    accuracy_axis.legend()
    accuracy_axis.bar_label(baseline_bars, fmt="%.3f", padding=3, fontsize=8)
    accuracy_axis.bar_label(adapted_bars, fmt="%.3f", padding=3, fontsize=8)

    speed_labels = ["Base Whisper", variant_label]
    speed_values = [
        float(baseline.get("aggregate_real_time_factor", 0.0)),
        float(adapted.get("aggregate_real_time_factor", 0.0)),
    ]
    speed_bars = speed_axis.bar(speed_labels, speed_values, color=["#4C78A8", "#F58518"])
    speed_axis.set_ylim(0, max(0.05, *speed_values) * 1.20)
    speed_axis.set_ylabel("Aggregate real-time factor")
    speed_axis.set_title("Batch evaluation speed (lower is better)")
    speed_axis.grid(axis="y", alpha=0.25)
    speed_axis.bar_label(speed_bars, fmt="%.3f", padding=3, fontsize=8)

    sample_count = int(baseline.get("sample_count", 0))
    fig.suptitle(f"Base Whisper vs. {method_label} | fixed test split, n={sample_count}")
    fig.tight_layout()
    fig.savefig(output_path, dpi=180, bbox_inches="tight")
    return output_path


def _set_font(run: Any, size: float, color: RGBColor, bold: bool = False) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in tokens.items():
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def _set_cell_margins(
    cell: Any,
    top: int = 80,
    start: int = 120,
    bottom: int = 80,
    end: int = 120,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table: Any, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def _shade_cell(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _shade_paragraph(paragraph: Any, fill: str) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    shading = paragraph_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        paragraph_properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _mark_header_row(row: Any) -> None:
    row_properties = row._tr.get_or_add_trPr()
    marker = row_properties.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        row_properties.append(marker)
    marker.set(qn("w:val"), "true")


def _page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    _set_font(run, 9, GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _add_metadata(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    label_run = paragraph.add_run(f"{label}: ")
    _set_font(label_run, 11, RGBColor(0, 0, 0), bold=True)
    value_run = paragraph.add_run(value)
    _set_font(value_run, 11, RGBColor(0, 0, 0))


def _format_metric(value: float, percent: bool = True) -> str:
    return f"{value:.1%}" if percent else f"{value:.3f}"


def _new_report_document() -> tuple[Document, Any]:
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _configure_styles(document)

    header = section.header.paragraphs[0]
    header.text = "AI Specialist | Manufacturing ASR Evaluation"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        _set_font(run, 9, GRAY)
    _page_number(section.footer.paragraphs[0])
    return document, section


def _add_report_title(
    document: Document,
    title: str,
    subtitle_text: str,
    owner: str,
    run_id: str,
    backend: str,
    model_repo: str,
    model_revision: str,
) -> None:
    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_before = Pt(16)
    title_paragraph.paragraph_format.space_after = Pt(4)
    title_run = title_paragraph.add_run(title)
    _set_font(title_run, 23, RGBColor(0, 0, 0), bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run(subtitle_text)
    _set_font(subtitle_run, 14, GRAY)

    _add_metadata(document, "Owner", owner)
    _add_metadata(document, "Run ID", run_id)
    _add_metadata(document, "Generated", datetime.now().astimezone().isoformat(timespec="seconds"))
    _add_metadata(document, "Backend", backend)
    _add_metadata(document, "Model", f"{model_repo}@{model_revision}")


def _add_adaptation_results_table(
    document: Document,
    metrics: dict[str, Any],
    variant_key: str,
    variant_label: str,
) -> None:
    baseline = metrics["baseline"]
    adapted = metrics[variant_key]
    table = document.add_table(rows=1, cols=4)
    _set_table_geometry(table, [3000, 2000, 2000, 2360])
    _mark_header_row(table.rows[0])
    for cell, label in zip(
        table.rows[0].cells,
        ["Metric", "Base Whisper", variant_label, f"{variant_label} - Base"],
        strict=True,
    ):
        _shade_cell(cell, LIGHT_GRAY_HEX)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(label)
        _set_font(run, 10, RGBColor(0, 0, 0), bold=True)

    rows = [
        ("WER (lower is better)", "wer", True),
        ("CER (lower is better)", "cer", True),
        ("Evaluation runtime (seconds)", "evaluation_runtime_seconds", False),
        ("Aggregate real-time factor", "aggregate_real_time_factor", False),
    ]
    for label, key, percent in rows:
        baseline_value = float(baseline.get(key, 0.0))
        adapted_value = float(adapted.get(key, 0.0))
        delta = adapted_value - baseline_value
        row = table.add_row().cells
        row[0].text = label
        row[1].text = _format_metric(baseline_value, percent)
        row[2].text = _format_metric(adapted_value, percent)
        row[3].text = f"{delta:+.1%}p" if percent else f"{delta:+.3f}"
        for index, cell in enumerate(row):
            cell.paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            for run in cell.paragraphs[0].runs:
                _set_font(run, 10, RGBColor(0, 0, 0))
        if delta < 0:
            for run in row[3].paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = DARK_BLUE
    _set_table_geometry(table, [3000, 2000, 2000, 2360])


def _build_lora_report(
    output_path: Path,
    title: str,
    run_id: str,
    owner: str,
    backend: str,
    model_repo: str,
    model_revision: str,
    metrics: dict[str, Any],
    chart_path: Path | None,
) -> Path:
    is_distillation = "distilled" in metrics
    variant_key = "distilled" if is_distillation else "lora"
    improvement_key = "distilled_improvement" if is_distillation else "lora_improvement"
    variant_label = "Distilled Student" if is_distillation else "Best LoRA"
    method_label = "knowledge distillation" if is_distillation else "LoRA"
    document, _ = _new_report_document()
    _add_report_title(
        document,
        title,
        f"고정 테스트셋 기반 Base Whisper·{variant_label} 비교 및 백데이터 기록",
        owner,
        run_id,
        backend,
        model_repo,
        model_revision,
    )

    baseline = metrics["baseline"]
    adapted = metrics[variant_key]
    reduction = float(metrics[improvement_key]["wer_absolute_reduction"])
    direction = "감소" if reduction >= 0 else "증가"
    confidence = str(metrics.get("evidence", {}).get("confidence", "low"))

    document.add_heading("1. Technical Summary", level=1)
    summary = document.add_paragraph()
    summary.paragraph_format.left_indent = Inches(0.12)
    summary.paragraph_format.right_indent = Inches(0.12)
    summary.paragraph_format.space_before = Pt(4)
    summary.paragraph_format.space_after = Pt(8)
    _shade_paragraph(summary, CALLOUT_HEX)
    summary.add_run(
        f"동일한 테스트 {int(baseline['sample_count'])}개에서 {variant_label}의 WER는 "
        f"{baseline['wer']:.1%}에서 {adapted['wer']:.1%}로 "
        f"{abs(reduction):.1%}p {direction}했습니다. "
        f"근거 신뢰도는 {confidence}이며, 공개 일반 한국어 소표본 결과이므로 제조 현장 성능으로 "
        "해석할 수 없습니다."
    )

    document.add_heading("2. Key Findings with Visual Evidence", level=1)
    _add_adaptation_results_table(document, metrics, variant_key, variant_label)
    if chart_path and chart_path.exists():
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_run = paragraph.add_run()
        picture_run.add_picture(str(chart_path), width=Inches(6.2))
        properties = picture_run._r.xpath(".//wp:docPr")
        if properties:
            properties[0].set(
                "descr",
                f"Grouped bars comparing base Whisper and {variant_label} WER, CER, and speed",
            )
            properties[0].set("title", f"Base Whisper versus {variant_label} comparison")
    document.add_paragraph(
        "WER와 CER는 같은 모델 리비전, 같은 디코딩 설정, 같은 테스트 행 순서에서 계산했습니다. "
        "속도는 전체 배치 평가 시간을 전체 음성 길이로 나눈 aggregate RTF이므로 단일 요청 "
        "지연시간과 동일하지 않습니다."
    )

    training = metrics.get("training", {})
    evidence = metrics.get("evidence", {})
    document.add_heading("3. Scope, Data, and Metric Definitions", level=1)
    document.add_paragraph(
        f"데이터: {evidence.get('dataset', 'public Korean speech sample')}; "
        f"train {int(training.get('train_samples', 0))}, validation "
        f"{int(training.get('validation_samples', 0))}, test {int(baseline['sample_count'])}."
    )
    for text in [
        "WER: 정규화된 단어열의 삽입·삭제·대치를 기준으로 한 오류율(낮을수록 좋음).",
        "CER: 정규화된 문자열의 오류율(낮을수록 좋음).",
        "Aggregate RTF: 전체 테스트 생성 시간 / 전체 음성 길이(낮을수록 빠름).",
        "Domain term recall: 테스트 정답에 등록 제조 용어가 있을 때만 적용하며, 없으면 N/A.",
    ]:
        document.add_paragraph(text, style="List Bullet")

    document.add_heading("4. Methodology", level=1)
    if is_distillation:
        document.add_paragraph(
            "고정된 Hugging Face 커밋에서 teacher와 student Whisper를 불러왔습니다. 정답 token의 "
            "교차엔트로피와 temperature-scaled teacher KL divergence를 결합해 student를 학습하고, "
            "학습 전후 student를 동일한 test split에서 비교했습니다."
        )
    else:
        document.add_paragraph(
            "고정된 Hugging Face 모델 커밋에서 Whisper processor와 base model을 불러왔습니다. "
            "LoRA는 q_proj·v_proj에 적용하고 validation WER가 가장 낮은 checkpoint를 "
            "자동 복원한 뒤 "
            "adapter를 저장했습니다. PEFT adapter 비활성화 상태를 Base Whisper, 활성화 상태를 "
            "Best LoRA로 두어 동일 test split을 연속 평가했습니다."
        )
    document.add_paragraph(
        f"학습 설정: max_steps={int(training.get('max_steps', 0))}, "
        f"best validation WER={float(training.get('best_validation_wer', 0.0)):.3f}, "
        f"seed={int(training.get('seed', 0))}."
    )

    document.add_heading("5. Limitations, Uncertainty, and Robustness", level=1)
    for text in [
        "테스트 표본이 작아 신뢰구간과 안정적인 일반화 오차를 제시하지 않았습니다.",
        "Zeroth-Korean은 일반 한국어 공개 데이터이며 제조 소음·화자·용어 분포를 대표하지 않습니다.",
        "단일 seed와 짧은 학습만 수행해 hyperparameter 및 반복 실행 민감도를 평가하지 않았습니다.",
        "예측과 지표는 자동 계산되었지만 transcript 품질과 보고서 결론은 아직 현업 검수 전입니다.",
    ]:
        document.add_paragraph(text, style="List Bullet")

    document.add_heading("6. Recommended Next Steps", level=1)
    for text in [
        "승인된 비식별 제조 음성으로 고정 test split과 제조 용어 평가셋을 구축합니다.",
        "3개 이상 seed와 더 긴 학습 설정으로 WER/CER 변동성과 과적합을 확인합니다.",
        "오류 유형·용어별 recall·소음 조건별 성능을 추가하고 사람이 예측 표본을 검수합니다.",
        "정확도와 함께 GPU 메모리, batch/streaming 지연시간, 운영 비용을 비교합니다.",
    ]:
        document.add_paragraph(text, style="List Number")

    document.add_heading("7. Further Questions", level=1)
    document.add_paragraph(
        "실제 현장에서 가장 중요한 음성 조건과 제조 용어는 무엇인지, 허용 가능한 WER·지연시간·비용 "
        "기준은 무엇인지, 데이터 사용과 보존 승인은 누가 담당하는지 확정해야 합니다."
    )

    document.core_properties.title = title
    document.core_properties.subject = f"AI Specialist Whisper {method_label} evaluation report"
    document.core_properties.author = owner
    document.save(output_path)
    return output_path


def build_report(
    output_path: Path,
    title: str,
    run_id: str,
    owner: str,
    backend: str,
    model_repo: str,
    model_revision: str,
    metrics: dict[str, Any],
    chart_path: Path | None,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if "lora" in metrics or "distilled" in metrics:
        return _build_lora_report(
            output_path,
            title,
            run_id,
            owner,
            backend,
            model_repo,
            model_revision,
            metrics,
            chart_path,
        )

    document, _ = _new_report_document()
    _add_report_title(
        document,
        title,
        "재현 가능한 ASR 실험·평가·백데이터 자동화 결과",
        owner,
        run_id,
        backend,
        model_repo,
        model_revision,
    )

    document.add_heading("1. Executive Summary", level=1)
    improvement = metrics["improvement"]
    summary = document.add_paragraph()
    summary.paragraph_format.left_indent = Inches(0.12)
    summary.paragraph_format.right_indent = Inches(0.12)
    summary.paragraph_format.space_before = Pt(4)
    summary.paragraph_format.space_after = Pt(8)
    _shade_paragraph(summary, CALLOUT_HEX)
    summary_text = (
        "도메인 용어 보정 적용 후 WER 절대 개선은 "
        f"{improvement['wer_absolute_reduction']:.1%}p입니다. "
    )
    if improvement.get("domain_term_recall_applicable"):
        summary_text += f"용어 재현율 개선은 {improvement['domain_term_recall_gain']:.1%}p입니다. "
    else:
        summary_text += "평가 문장에 사전의 제조 용어가 없어 용어 재현율은 N/A입니다. "
    summary.add_run(
        summary_text
        + "본 결과는 실행별 설정과 원본 예측이 함께 보존되어 재현 및 감사가 가능합니다."
    )

    document.add_heading("2. Evaluation Results", level=1)
    table = document.add_table(rows=1, cols=4)
    _set_table_geometry(table, [3000, 2000, 2000, 2360])
    _mark_header_row(table.rows[0])
    headers = ["Metric", "Baseline", "Corrected", "Absolute change"]
    for cell, label in zip(table.rows[0].cells, headers, strict=True):
        _shade_cell(cell, LIGHT_GRAY_HEX)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(label)
        _set_font(run, 10, RGBColor(0, 0, 0), bold=True)

    rows = [
        ("WER (lower is better)", "wer", -1),
        ("CER (lower is better)", "cer", -1),
        ("Domain Term Precision", "domain_term_precision", 1),
        ("Domain Term Recall", "domain_term_recall", 1),
        ("Domain Term F1-score", "domain_term_f1", 1),
        ("Mean Latency (seconds)", "mean_latency_seconds", 0),
        ("Mean Real-time Factor", "mean_real_time_factor", 0),
    ]
    for label, key, direction in rows:
        baseline_value = float(metrics["baseline"][key])
        corrected_value = float(metrics["corrected"][key])
        row = table.add_row().cells
        row[0].text = label
        percent = key in {
            "wer",
            "cer",
            "domain_term_precision",
            "domain_term_recall",
            "domain_term_f1",
        }
        delta = corrected_value - baseline_value
        applicable = not (
            key in {"domain_term_precision", "domain_term_recall", "domain_term_f1"}
            and not metrics["baseline"].get(
                "domain_term_recall_applicable"
                if key != "domain_term_precision"
                else "domain_term_precision_applicable"
            )
        )
        if applicable:
            row[1].text = _format_metric(baseline_value, percent)
            row[2].text = _format_metric(corrected_value, percent)
            row[3].text = f"{delta:+.1%}p" if percent else f"{delta:+.3f}"
        else:
            row[1].text = "N/A"
            row[2].text = "N/A"
            row[3].text = "N/A"
        for index, cell in enumerate(row):
            cell.paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            for run in cell.paragraphs[0].runs:
                _set_font(run, 10, RGBColor(0, 0, 0))
        if (
            applicable
            and direction
            and ((direction < 0 and delta < 0) or (direction > 0 and delta > 0))
        ):
            for run in row[3].paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = DARK_BLUE
    _set_table_geometry(table, [3000, 2000, 2000, 2360])

    if chart_path and chart_path.exists():
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_run = paragraph.add_run()
        picture_run.add_picture(str(chart_path), width=Inches(6.1))
        document_properties = picture_run._r.xpath(".//wp:docPr")
        if document_properties:
            document_properties[0].set(
                "descr",
                "Bar chart comparing baseline and corrected WER, CER, and domain term recall",
            )
            document_properties[0].set("title", "ASR evaluation metric comparison")

    quality_gate = metrics.get("quality_gate")
    next_section = 3
    if isinstance(quality_gate, dict) and quality_gate.get("enabled"):
        document.add_heading("3. Manufacturing Quality Gate", level=1)
        target_table = document.add_table(rows=1, cols=5)
        target_table.style = "Table Grid"
        target_headers = ["Priority", "Metric", "Target", "Observed", "Status"]
        for cell, label in zip(target_table.rows[0].cells, target_headers, strict=True):
            cell.text = label
        target_rows = [
            (
                "1",
                "Domain Term Recall",
                f">= {float(quality_gate['targets']['minimum_domain_term_recall']):.1%}",
                f"{float(quality_gate['observed']['domain_term_recall']):.1%}",
                quality_gate["checks"]["domain_term_recall"]["passed"],
            ),
            (
                "2",
                "CER",
                f"<= {float(quality_gate['targets']['maximum_cer']):.1%}",
                f"{float(quality_gate['observed']['cer']):.1%}",
                quality_gate["checks"]["cer"]["passed"],
            ),
            (
                "3",
                "WER",
                f"<= {float(quality_gate['targets']['maximum_wer']):.1%}",
                f"{float(quality_gate['observed']['wer']):.1%}",
                quality_gate["checks"]["wer"]["passed"],
            ),
        ]
        for values in target_rows:
            row = target_table.add_row().cells
            for cell, value in zip(row[:4], values[:-1], strict=True):
                cell.text = str(value)
            row[4].text = "PASS" if values[-1] else "FAIL"
        overall_label = "PASS" if quality_gate["overall_pass"] else "FAIL"
        document.add_paragraph(
            f"Overall target status: {overall_label}. "
            "세 지표를 모두 만족해야 현업 정확도 목표 통과 후보로 분류합니다."
        )
        next_section = 4

    document.add_heading(f"{next_section}. Interpretation and Limits", level=1)
    if backend == "fixture":
        document.add_paragraph(
            "Fixture 실행은 자동화 구조와 계산 로직을 검증하기 위한 합성 결과입니다. "
            "실제 성능 주장은 현업 담당자가 검수한 음성·전사 데이터의 고정된 test split에서만 "
            "작성해야 합니다."
        )
    else:
        document.add_paragraph(
            "공개 일반 한국어 데이터 실행은 모델·데이터·보고서 경로의 기술 검증입니다. 제조 현장 "
            "성능 주장은 현업 담당자가 검수한 제조 음성·전사 데이터의 고정된 test split에서만 "
            "작성해야 합니다."
        )
    document.add_paragraph(
        "도메인 용어 보정은 사전 기반 후처리이므로 문맥상 잘못된 치환 가능성을 별도 오류 분석에서 "
        "확인해야 합니다. 모델 선정은 WER뿐 아니라 용어 정확도, 지연시간, 메모리, 보안 조건을 함께 "
        "평가해야 합니다."
    )

    document.add_heading(f"{next_section + 1}. Next Review Gate", level=1)
    document.add_paragraph(
        "다음 단계는 실제 데이터의 사용 승인과 transcript 표본 검수를 완료한 뒤 동일 "
        "파이프라인으로 "
        "Baseline을 측정하는 것입니다. Baseline 결과가 확정된 이후에만 LoRA fine-tuning과 INT8 "
        "비교 실험을 활성화합니다."
    )

    document.core_properties.title = title
    document.core_properties.subject = "AI Specialist manufacturing ASR experiment report"
    document.core_properties.author = owner
    document.save(output_path)
    return output_path
