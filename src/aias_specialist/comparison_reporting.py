from __future__ import annotations

import os
import tempfile
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


def create_comparison_chart(frame: pd.DataFrame, output_path: Path, title: str) -> Path:
    matplotlib_cache = Path(tempfile.gettempdir()) / "aias-matplotlib"
    matplotlib_cache.mkdir(parents=True, exist_ok=True)
    os.environ.setdefault("MPLCONFIGDIR", str(matplotlib_cache))
    from matplotlib.figure import Figure

    completed = frame.loc[frame["status"] == "completed"].copy()
    if completed.empty:
        raise ValueError("A comparison chart requires at least one completed member")
    labels = completed["member_id"].astype(str).tolist()
    positions = list(range(len(labels)))
    width = 0.36
    fig = Figure(figsize=(max(8.0, len(labels) * 1.5), 4.8))
    accuracy_axis, speed_axis = fig.subplots(1, 2)

    accuracy_axis.bar(
        [position - width / 2 for position in positions],
        completed["wer"].astype(float),
        width,
        label="WER",
    )
    accuracy_axis.bar(
        [position + width / 2 for position in positions],
        completed["cer"].astype(float),
        width,
        label="CER",
    )
    accuracy_axis.set_xticks(positions, labels, rotation=25, ha="right")
    accuracy_axis.set_ylabel("Error rate (lower is better)")
    accuracy_axis.set_title("Recognition accuracy")
    accuracy_axis.grid(axis="y", alpha=0.25)
    accuracy_axis.legend()

    speed_axis.scatter(
        completed["aggregate_real_time_factor"].astype(float),
        completed["cer"].astype(float),
        s=70,
    )
    for row in completed.itertuples(index=False):
        speed_axis.annotate(
            str(row.member_id),
            (float(row.aggregate_real_time_factor), float(row.cer)),
            xytext=(4, 4),
            textcoords="offset points",
            fontsize=8,
        )
    speed_axis.set_xlabel("Aggregate real-time factor (lower is faster)")
    speed_axis.set_ylabel("CER (lower is better)")
    speed_axis.set_title("Accuracy-speed trade-off")
    speed_axis.grid(alpha=0.25)
    fig.suptitle(title)
    fig.tight_layout()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, dpi=180, bbox_inches="tight")
    return output_path


def build_comparison_report(
    frame: pd.DataFrame,
    output_path: Path,
    *,
    title: str,
    group_id: str,
    group_kind: str,
    evaluation_split: str,
    owner: str,
    chart_path: Path | None,
    reference_member: str | None = None,
) -> Path:
    completed = frame.loc[frame["status"] == "completed"].copy()
    if completed.empty:
        raise ValueError("A comparison report requires at least one completed member")
    if "rank" in completed and completed["rank"].notna().any():
        ranked = completed.sort_values("rank").reset_index(drop=True)
    else:
        ranked = completed.sort_values(
            ["cer", "wer", "aggregate_real_time_factor"],
            ascending=[True, True, True],
        ).reset_index(drop=True)
    best = ranked.iloc[0]
    targets_active = (
        "quality_targets_enabled" in ranked
        and ranked["quality_targets_enabled"]
        .map(lambda value: str(value).strip().lower() == "true")
        .any()
    )

    document = Document()
    document.add_heading(title, level=0)
    document.add_paragraph(f"Experiment group: {group_id}")
    document.add_paragraph(f"Owner: {owner}")
    document.add_paragraph(f"Evaluation split: {evaluation_split}")

    document.add_heading("1. Technical Summary", level=1)
    kind_label = "모델" if group_kind == "benchmark" else "양자화 단계"
    document.add_paragraph(
        f"동일한 {evaluation_split} 데이터와 디코딩 조건에서 {len(completed)}개 {kind_label}를 "
        f"비교했습니다. 목표 우선순위 기반 1위는 {best['member_id']}이며 "
        f"용어 Recall={float(best['domain_term_recall']):.3f}, "
        f"CER={float(best['cer']):.3f}, "
        f"WER={float(best['wer']):.3f}, aggregate RTF="
        f"{float(best['aggregate_real_time_factor']):.3f}입니다."
    )

    document.add_heading("2. Comparison Results", level=1)
    table = document.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = [
        "Rank",
        "Candidate",
        "WER",
        "CER",
        "Term recall",
        "RTF",
        "GPU MB",
        "Model MB",
    ]
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = header
    for rank, row in ranked.iterrows():
        cells = table.add_row().cells
        values = [
            str(rank + 1),
            str(row["member_id"]),
            f"{float(row['wer']):.3f}",
            f"{float(row['cer']):.3f}",
            f"{float(row['domain_term_recall']):.3f}",
            f"{float(row['aggregate_real_time_factor']):.3f}",
            f"{float(row['peak_gpu_memory_mb']):.0f}",
            f"{float(row['model_size_bytes']) / (1024 * 1024):.1f}",
        ]
        for cell, value in zip(cells, values, strict=True):
            cell.text = value

    if targets_active:
        document.add_heading("3. Manufacturing Quality Targets", level=1)
        target_table = document.add_table(rows=1, cols=7)
        target_table.style = "Table Grid"
        target_headers = [
            "Candidate",
            "Overall",
            "Recall gap",
            "CER gap",
            "WER gap",
            "Recall pass",
            "CER/WER pass",
        ]
        for cell, header in zip(target_table.rows[0].cells, target_headers, strict=True):
            cell.text = header
        for row in ranked.itertuples(index=False):
            cells = target_table.add_row().cells
            values = [
                str(row.member_id),
                "PASS" if bool(row.quality_target_pass) else "FAIL",
                f"{float(row.domain_term_recall_gap):.1%}p",
                f"{float(row.cer_gap):.1%}p",
                f"{float(row.wer_gap):.1%}p",
                "PASS" if bool(row.domain_term_recall_target_pass) else "FAIL",
                ("PASS" if bool(row.cer_target_pass) and bool(row.wer_target_pass) else "FAIL"),
            ]
            for cell, value in zip(cells, values, strict=True):
                cell.text = value

    if chart_path and chart_path.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(chart_path), width=Inches(6.3))

    if group_kind == "quantization" and reference_member:
        tradeoff_number = 4 if targets_active else 3
        document.add_heading(f"{tradeoff_number}. Quantization Trade-off", level=1)
        document.add_paragraph(
            f"정확도·속도·크기 변화는 {reference_member} 단계를 기준으로 계산했습니다. "
            "CER/WER 증가는 정확도 저하, RTF speedup은 1보다 클수록 속도 개선을 뜻합니다."
        )
        tradeoff = document.add_table(rows=1, cols=6)
        tradeoff.style = "Table Grid"
        headers = [
            "Variant",
            "CER Δ",
            "WER Δ",
            "Term recall Δ",
            "RTF speedup",
            "Size reduction",
        ]
        for cell, header in zip(tradeoff.rows[0].cells, headers, strict=True):
            cell.text = header
        for row in ranked.itertuples(index=False):
            cells = tradeoff.add_row().cells
            values = [
                str(row.member_id),
                f"{float(row.cer_delta_vs_reference):+.3f}",
                f"{float(row.wer_delta_vs_reference):+.3f}",
                f"{float(row.term_recall_delta_vs_reference):+.3f}",
                f"{float(row.rtf_speedup_vs_reference):.2f}x",
                f"{float(row.model_size_reduction_ratio):.1%}",
            ]
            for cell, value in zip(cells, values, strict=True):
                cell.text = value

    methodology_number = 4 if group_kind == "quantization" else 3
    if targets_active:
        methodology_number += 1
    document.add_heading(f"{methodology_number}. Methodology and Traceability", level=1)
    document.add_paragraph(
        "모든 후보는 독립적인 run_id로 실행되며 설정 snapshot, 환경, 입력 manifest, 원본 예측, "
        "지표와 Word 보고서를 보존합니다. 비교표는 원본 ASR 출력의 baseline 지표로 작성되며 "
        "도메인 용어 후처리 결과와 혼합하지 않습니다."
    )

    document.add_heading(
        f"{methodology_number + 1}. Limitations and Human Review Gate",
        level=1,
    )
    for text in [
        "Validation 결과는 모델과 양자화 선택용이며 최종 성능 주장은 고정 Test 실행으로 "
        "확인해야 합니다.",
        "녹음 정답, 제조 용어, 개인정보·외부 반출 승인은 현업 담당자의 검수가 필요합니다.",
        "최저 CER 후보가 운영상 최선이라고 자동 확정하지 않으며 속도·메모리·모델 크기를 "
        "함께 검토합니다.",
        "공개 또는 fixture 데이터 결과는 제조 현장 성능 근거로 사용할 수 없습니다.",
    ]:
        document.add_paragraph(text, style="List Bullet")

    document.add_heading(f"{methodology_number + 2}. Selection Record", level=1)
    document.add_paragraph(
        "추천 순위는 자동 생성되지만 최종 선택은 별도의 select 명령으로 담당자와 선택 이유를 "
        "기록해야 합니다."
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = title
    document.core_properties.subject = "ASR model and quantization comparison"
    document.core_properties.author = owner
    document.save(output_path)
    return output_path
